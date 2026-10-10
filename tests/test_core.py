from dataclasses import replace
from pathlib import Path
import json
import pytest
from openpyxl import Workbook,load_workbook
from pypdf import PdfWriter
from vat_core import Period,load_clients,build_jobs,validate_artifact,Journal,biz_number
from prepare_input import prepare


def workbook(path, rows=None, extra=None):
    w=Workbook();s=w.active
    s.append(['NO','사업자명','대표이사','사업자등록번호','과세기간']+list((extra or {}).keys()))
    for row in rows or [[1,'테스트업체','','123-45-67890','2기예정']]:
        s.append(row+list((extra or {}).values()))
    w.save(path);return path


@pytest.mark.parametrize('label,scope,start,end,quarters,months',[
 ('1기예정',None,'2026-01-01','2026-03-31',(1,),(1,2,3)),
 ('1기확정',None,'2026-01-01','2026-06-30',(1,2),(1,2,3,4,5,6)),
 ('2기예정',None,'2026-07-01','2026-09-30',(3,),(7,8,9)),
 ('2기확정',None,'2026-07-01','2026-12-31',(3,4),(7,8,9,10,11,12)),
 ('2기확정','분기','2026-10-01','2026-12-31',(4,),(10,11,12)),
 ('1기확정','분기','2026-04-01','2026-06-30',(2,),(4,5,6)),
])
def test_tax_ranges(label,scope,start,end,quarters,months):
    p=Period.parse(label,2026,scope)
    assert (p.start,p.end,p.quarters,p.months)==(start,end,quarters,months)


@pytest.mark.parametrize('period,year,scope',[('2기',2026,None),('3기예정',2026,None),('2기예정',None,None),
 ('2기예정','2026.0',None),('2기예정',2026,'반기'),('1기확정',2026,'전체')])
def test_reject_ambiguous_period(period,year,scope):
    with pytest.raises(ValueError):Period.parse(period,year,scope)


@pytest.mark.parametrize('value',['123-45-67890','1234567890',1234567890,1234567890.0])
def test_biz_number(value):assert biz_number(value)=='1234567890'


@pytest.mark.parametrize('value',['123456789','12x4567890','12345678901','1234567890.0',None])
def test_reject_lost_or_invalid_biz(value):
    with pytest.raises(ValueError):biz_number(value)


def test_year_required_and_fallback(tmp_path):
    p=workbook(tmp_path/'a.xlsx')
    with pytest.raises(ValueError,match='과세연도'):load_clients(p)
    assert load_clients(p,2026)[0].period.year==2026


def test_input_formula_duplicate_headers_and_rows(tmp_path):
    p=workbook(tmp_path/'formula.xlsx',extra={'과세연도':'=2026'})
    with pytest.raises(ValueError,match='수식'):load_clients(p)
    p=workbook(tmp_path/'duplicate.xlsx',rows=[[1,'A','','1234567890','2기예정'],[2,'B','','1234567890','2기예정']],extra={'과세연도':2026})
    with pytest.raises(ValueError,match='중복'):load_clients(p)
    w=load_workbook(p);w.active.cell(1,2,'NO');w.save(p)
    with pytest.raises(ValueError,match='헤더'):load_clients(p)


def test_plan_has_all_categories_and_four_invoice_variants(tmp_path):
    c=load_clients(workbook(tmp_path/'a.xlsx',extra={'과세연도':2026}))
    jobs=build_jobs(c)
    assert len(jobs)==13
    assert {j.service for j in jobs}==set(range(1,8))
    invoices=[j for j in jobs if j.service==2]
    assert {j.variant for j in invoices}=={'전자세금계산서_매출','전자세금계산서_매입','전자계산서_매출','전자계산서_매입'}
    assert all((j.filters['시작일'],j.filters['종료일'])==('2026-07-01','2026-09-30') for j in invoices)
    assert [j.filters['조회월'] for j in jobs if j.service==4]==['7월','8월','9월']
    assert {j.variant for j in jobs if j.service==7}=={'내국신용장','구매확인서'}


def test_final_filing_half_year_and_special_quarters(tmp_path):
    c=load_clients(workbook(tmp_path/'a.xlsx',rows=[[1,'A','','1234567890','2기확정']],extra={'과세연도':2026}))
    jobs=build_jobs(c)
    assert [j.filters['적용분기'] for j in jobs if j.service==5]==['3분기','4분기']
    assert [j.filters['조회월'] for j in jobs if j.service==4]==['7월','8월','9월','10월','11월','12월']
    export=next(j for j in jobs if j.service==6)
    assert (export.filters['조회단위'],export.filters['조회기간'])==('반기','하반기')


def test_reject_trade_final_quarter_only(tmp_path):
    c=load_clients(workbook(tmp_path/'a.xlsx',rows=[[1,'A','','1234567890','2기확정']],extra={'과세연도':2026,'조회범위':'분기'}))
    with pytest.raises(ValueError,match='7번'):build_jobs(c)


def test_skip_disabled_row_and_select_services(tmp_path):
    p=workbook(tmp_path/'a.xlsx',extra={'과세연도':2026,'조회항목':'2,7','실행여부':'Y'})
    assert {j.service for j in build_jobs(load_clients(p))}=={2,7}
    w=load_workbook(p);w.active.cell(2,8,'N');w.save(p)
    with pytest.raises(ValueError,match='거래처'):load_clients(p)


def test_template_preserves_source_and_adds_defaults(tmp_path):
    source=workbook(tmp_path/'source.xlsx');before=source.read_bytes()
    target=tmp_path/'expanded.xlsx';prepare(source,target,2026)
    assert source.read_bytes()==before
    assert len(build_jobs(load_clients(target)))==13
    with pytest.raises(FileExistsError):prepare(source,target)


def pdf(path):
    w=PdfWriter();w.add_blank_page(width=100,height=100)
    with path.open('wb') as f:w.write(f)
    return path


def test_file_validation_and_journal_resume(tmp_path):
    job=build_jobs(load_clients(workbook(tmp_path/'a.xlsx',extra={'과세연도':2026})))[0]
    path=pdf(tmp_path/'report.pdf');info=validate_artifact(path)
    assert info['pages']==1 and not info['text_present']
    journal=Journal(tmp_path/'history.jsonl')
    journal.record(job,'saved',[{'path':str(path),**info}],True)
    assert Journal(journal.path).reusable(job)
    # 순서/회사명 변경으로 동일 조회를 재처리하지 않습니다.
    assert journal.reusable(replace(job,client=replace(job.client,row=10,company='변경명')))
    path.write_bytes(b'%PDF-invalid')
    assert not journal.reusable(job)
    path=pdf(path);journal.record(job,'saved',[{'path':str(path),**validate_artifact(path)}],False)
    assert not journal.reusable(job)


@pytest.mark.parametrize('name,data',[('a.pdf',b'fake'),('a.xlsx',b''),('a.xls',b'<html>login</html>'),('a.exe',b'not data')])
def test_artifacts_reject_wrong_files(tmp_path,name,data):
    p=tmp_path/name;p.write_bytes(data)
    with pytest.raises(ValueError):validate_artifact(p)


def test_incomplete_journal_recoverable(tmp_path):
    job=build_jobs(load_clients(workbook(tmp_path/'a.xlsx',extra={'과세연도':2026})))[0]
    path=tmp_path/'journal.jsonl';path.write_text('{"job_key":',encoding='utf-8')
    journal=Journal(path);journal.record(job,'failed',[],False)
    assert len(Journal(path).latest)==1
    assert list(tmp_path.glob('journal.jsonl.incomplete-*'))


def test_interior_corrupt_journal_rejected(tmp_path):
    p=tmp_path/'j';p.write_text('bad\n{}\n')
    with pytest.raises(ValueError,match='손상'):Journal(p)
