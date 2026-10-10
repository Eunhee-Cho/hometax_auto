"""1~7 홈택스 자료 수집. 로그인/권한은 사용자가 완료하며 신고·납부하지 않습니다.

기본 auto 모드: 로그인 후 메뉴·거래처·기간 설정과 저장을 자동 진행합니다.
assisted 모드에서만 수동 Enter 확인을 사용합니다.
strict 모드: 선택자/대상 검증이 안 되면 실패로 기록합니다.
--validate-only 는 브라우저를 열지 않습니다.
"""
from __future__ import annotations
import argparse
from dataclasses import replace
import json
import os
import re
import shutil
import subprocess
import sys
import time
import urllib.request
from pathlib import Path
from urllib.parse import urlparse

# 직접 실행할 때만 누락된 필수 패키지를 설치합니다. 다른 소스 파일은 필요 없습니다.
def ensure_dependencies():
    import importlib.util
    if sys.version_info < (3, 11):
        raise RuntimeError('Python 3.11 이상이 필요합니다. python --version으로 확인하세요.')
    packages = {'openpyxl': 'openpyxl==3.1.5', 'playwright': 'playwright==1.62.0', 'pypdf': 'pypdf==6.20.0'}
    missing = [pin for module, pin in packages.items() if importlib.util.find_spec(module) is None]
    if missing:
        print('첫 실행: 필요한 패키지를 설치합니다. 인터넷 연결이 필요합니다.', flush=True)
        print('  ' + ', '.join(missing), flush=True)
        try:
            subprocess.run([sys.executable, '-m', 'pip', 'install', '--disable-pip-version-check',
                            '--retries', '1', '--timeout', '30', *missing], check=True, timeout=300)
        except (subprocess.SubprocessError, OSError) as exc:
            raise RuntimeError('패키지 설치 실패. 인터넷/pip 설정 확인 후 같은 명령을 다시 실행하세요.') from exc


if __name__ == '__main__':
    try:
        ensure_dependencies()
    except Exception as exc:
        print(f'실행 중단: {exc}', file=sys.stderr, flush=True)
        raise SystemExit(2)


# 입력·기간·작업계획·파일검증을 이 한 파일에 포함합니다.
import calendar
import hashlib
import json
import re
from dataclasses import dataclass, asdict
from datetime import date, datetime
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

SERVICES = {
    1: '부가가치세 신고도움서비스',
    2: '전자(세금)계산서 합계표조회',
    3: '신용카드 판매(결제)대행 매출자료 조회',
    4: '현금영수증 매출총액 조회',
    5: '부가가치세 매입자납부 특례 조회',
    6: '수출실적명세서 조회',
    7: '내국신용장 구매확인서 전자발급명세서 조회',
}
REQUIRED = ('사업자명', '사업자등록번호', '과세기간')


def clean(value: Any) -> str:
    return '' if value is None else str(value).strip()


def choices(value: Any, allowed: tuple[str, ...], default: str) -> tuple[str, ...]:
    text = clean(value) or default
    if text == '모두':
        return allowed
    parts = tuple(x.strip() for x in re.split('[,;|]', text) if x.strip())
    if not parts or len(parts) != len(set(parts)) or any(p not in allowed for p in parts):
        raise ValueError(f'허용값 {allowed} 또는 모두: {text}')
    return parts


def biz_number(value: Any) -> str:
    # 숫자로 저장된 앞자리 0은 원본 소실 여부를 판단할 수 없으므로 추측하지 않습니다.
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r'[-\s]', '', clean(value))
    if not re.fullmatch(r'[0-9]{10}', text):
        raise ValueError('사업자등록번호는 10자리 숫자(하이픈 허용)여야 합니다. 셀을 텍스트로 저장하세요.')
    return text


@dataclass(frozen=True)
class Period:
    year: int
    term: int
    filing: str
    scope: str

    @classmethod
    def parse(cls, text: Any, year: Any, scope: Any = None) -> 'Period':
        compact = re.sub(r'\s+', '', clean(text))
        m = re.fullmatch(r'([12])기(예정|확정)', compact)
        if not m:
            raise ValueError('과세기간은 1기예정/1기확정/2기예정/2기확정 중 하나여야 합니다.')
        if not re.fullmatch(r'[0-9]{4}', clean(year)):
            raise ValueError('과세연도 4자리가 필요합니다. 엑셀에 추가하거나 --year로 지정하세요.')
        y = int(clean(year))
        if not 2000 <= y <= 2100:
            raise ValueError('과세연도는 2000~2100 범위여야 합니다.')
        s = clean(scope) or ('분기' if m[2] == '예정' else '반기')
        if s not in ('분기', '반기'):
            raise ValueError('조회범위는 분기 또는 반기여야 합니다.')
        if m[2] == '예정' and s != '분기':
            raise ValueError('예정신고 조회범위는 분기여야 합니다.')
        return cls(y, int(m[1]), m[2], s)

    @property
    def quarters(self) -> tuple[int, ...]:
        base = 1 if self.term == 1 else 3
        if self.filing == '예정':
            return (base,)
        return (base, base + 1) if self.scope == '반기' else (base + 1,)

    @property
    def months(self) -> tuple[int, ...]:
        return tuple(m for q in self.quarters for m in range(q * 3 - 2, q * 3 + 1))

    @property
    def start(self) -> str:
        return date(self.year, self.months[0], 1).isoformat()

    @property
    def end(self) -> str:
        last = self.months[-1]
        return date(self.year, last, calendar.monthrange(self.year, last)[1]).isoformat()

    @property
    def label(self) -> str:
        return f'{self.year}_{self.term}기{self.filing}_{self.scope}'


@dataclass(frozen=True)
class Client:
    row: int
    company: str
    biz: str
    period: Period
    services: tuple[int, ...]
    invoice_types: tuple[str, ...]
    directions: tuple[str, ...]
    documents: tuple[str, ...]
    special_tabs: tuple[str, ...]
    goods: str
    special_direction: str
    formats: tuple[str, ...]


@dataclass(frozen=True)
class Job:
    client: Client
    service: int
    variant: str
    filters: dict[str, str]

    @property
    def key(self) -> str:
        raw = json.dumps({'biz':self.client.biz, 'period':asdict(self.client.period),
                          'formats':self.client.formats, 'service':self.service,
                          'variant':self.variant, 'filters':self.filters}, sort_keys=True, ensure_ascii=False)
        return hashlib.sha256(raw.encode()).hexdigest()[:24]

    @property
    def stem(self) -> str:
        # 전체 사업자번호를 파일명/일반 로그에 노출하지 않습니다.
        raw = f'{self.client.company}_{self.client.biz[-4:]}_{self.client.period.label}_{self.service:02d}_{SERVICES[self.service]}_{self.variant}'
        name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', raw).strip(' .')
        return name[:150] + '_' + self.key[:8]


def load_clients(path: Path, fallback_year: int | None = None, sheet: str | None = None) -> list[Client]:
    w = load_workbook(path, data_only=False, read_only=True)
    try:
        s = w[sheet] if sheet else w.active
        rows = iter(s.iter_rows())
        headers = [clean(c.value) for c in next(rows)]
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise ValueError('엑셀 헤더가 중복됩니다.')
        for key in REQUIRED:
            if key not in headers:
                raise ValueError(f'필수 열 없음: {key}')
        clients: list[Client] = []
        seen = set()
        for rownum, row in enumerate(rows, 2):
            if not any(c.value is not None for c in row):
                continue
            if any(c.data_type == 'f' for c in row):
                raise ValueError(f'{rownum}행: 수식 대신 확정된 입력값을 사용하세요.')
            data = dict(zip(headers, (c.value for c in row)))
            if clean(data.get('실행여부')).upper() in ('N', '아니오', '0'):
                continue
            try:
                company = clean(data.get('사업자명'))
                if not company:
                    raise ValueError('사업자명 누락')
                p = Period.parse(data.get('과세기간'), data.get('과세연도') or fallback_year, data.get('조회범위'))
                svc = clean(data.get('조회항목')) or '1,2,3,4,5,6,7'
                services = tuple(int(x.strip()) for x in re.split('[,;|]', svc))
                if not services or len(services) != len(set(services)) or any(x not in SERVICES for x in services):
                    raise ValueError('조회항목은 1~7 숫자를 쉼표로 구분하세요.')
                goods = clean(data.get('특례품목')) or '전체'
                sd = clean(data.get('특례매입매출구분')) or '전체'
                if sd not in ('전체', '매입', '매출'):
                    raise ValueError('특례매입매출구분은 전체/매입/매출')
                client = Client(rownum, company, biz_number(data.get('사업자등록번호')), p, services,
                    choices(data.get('계산서종류'), ('전자세금계산서','전자계산서'), '모두'),
                    choices(data.get('매출매입구분'), ('매출','매입'), '모두'),
                    choices(data.get('무역서류종류'), ('내국신용장','구매확인서'), '모두'),
                    choices(data.get('특례조회탭'), ('거래내역조회','전용계좌조회','국고입금내역','기납부세액'), '거래내역조회'),
                    goods, sd, choices(data.get('저장형식'), ('PDF','원본'), '모두'))
                identity = (client.biz, p)
                if identity in seen:
                    raise ValueError('동일 사업자·과세기간이 중복됩니다. 조회항목을 한 행에 통합하세요.')
                seen.add(identity)
                clients.append(client)
            except (ValueError, TypeError) as exc:
                raise ValueError(f'{rownum}행: {exc}') from exc
        if not clients:
            raise ValueError('실행할 거래처가 없습니다.')
        return clients
    finally:
        w.close()


def build_jobs(clients: list[Client]) -> list[Job]:
    jobs = []
    for c in clients:
        p = c.period
        base = {'사업자등록번호':c.biz, '과세연도':str(p.year)}
        for service in c.services:
            if service == 1:
                jobs.append(Job(c,1,'전체보고서',base | {'기수':f'{p.term}기','신고구분':p.filing}))
            elif service == 2:
                for typ in c.invoice_types:
                    for direction in c.directions:
                        jobs.append(Job(c,2,f'{typ}_{direction}',base | {'분류':typ,'구분':direction,
                            '작성일자':'일자별','시작일':p.start,'종료일':p.end}))
            elif service == 3:
                jobs.append(Job(c,3,'매출',base | {'시작분기':f'{min(p.quarters)}분기','종료분기':f'{max(p.quarters)}분기'}))
            elif service == 4:
                # 화면은 조회연도 + 월 선택. 전체를 선택하면 과세기간 외 자료가 섞입니다.
                for month in p.months:
                    jobs.append(Job(c,4,f'{month:02d}월',base | {'조회월':f'{month}월'}))
            elif service == 5:
                for tab in c.special_tabs:
                    for quarter in p.quarters:
                        jobs.append(Job(c,5,f'{tab}_{quarter}분기',base | {'특례조회탭':tab,
                            '적용분기':f'{quarter}분기','품목구분':c.goods,'매입매출구분':c.special_direction}))
            elif service == 6:
                scope = '반기' if p.filing == '확정' and p.scope == '반기' else '분기'
                selection = ('상반기' if p.term==1 else '하반기') if scope=='반기' else f'{p.quarters[0]}분기'
                jobs.append(Job(c,6,'수출실적',base | {'조회단위':scope,'조회기간':selection}))
            elif service == 7:
                if p.filing == '확정' and p.scope == '분기':
                    # 사이트 구매공급일의 확정 구분은 반기이며 4~6/10~12월만 따로 요청할 수 없습니다.
                    raise ValueError(f'{c.row}행: 7번은 분기만 조회하는 확정신고를 지원하지 않습니다. 조회범위=반기로 하거나 7번을 제외하세요.')
                for doc in c.documents:
                    jobs.append(Job(c,7,doc,base | {'서류구분':doc,'기수':f'{p.term}기','신고구분':p.filing}))
    return jobs


def validate_artifact(path: Path) -> dict:
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError('다운로드 파일이 없거나 비어 있습니다.')
    head = path.read_bytes()[:1024].lstrip()
    if head.lower().startswith((b'<!doctype html', b'<html')):
        raise ValueError('자료 대신 HTML/로그인 페이지가 다운로드되었습니다.')
    suffix = path.suffix.lower()
    info = {'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
    if suffix == '.pdf':
        if not head.startswith(b'%PDF-'):
            raise ValueError('PDF 헤더가 없습니다.')
        from pypdf import PdfReader
        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted or not len(reader.pages):
            raise ValueError('PDF가 암호화되었거나 페이지가 없습니다.')
        info['pages'] = len(reader.pages)
        # 이미지 PDF에는 텍스트가 없을 수 있어 텍스트 일치로 성공을 꾸미지 않습니다.
        info['text_present'] = bool(''.join(p.extract_text() or '' for p in reader.pages).strip())
    elif suffix == '.xlsx':
        wb = load_workbook(path, read_only=True)
        try:
            if not wb.sheetnames:
                raise ValueError('엑셀 시트가 없습니다.')
        finally:
            wb.close()
    elif suffix == '.xls':
        if not (head.startswith(bytes.fromhex('D0CF11E0A1B11AE1')) or b'<table' in head.lower() or b'<xml' in head.lower()):
            raise ValueError('알려진 XLS 형식이 아닙니다.')
        info['structure_checked'] = 'header_only'
    elif suffix not in ('.csv','.txt','.zip'):
        raise ValueError(f'예상하지 못한 확장자: {suffix}')
    elif suffix == '.zip':
        import zipfile
        with zipfile.ZipFile(path) as z:
            if not z.namelist() or z.testzip() is not None:
                raise ValueError('손상/빈 ZIP')
    return info


class Journal:
    """증분 JSONL. 개인정보가 포함되는 로컬 파일이며 외부 전송하지 않습니다."""
    def __init__(self, path: Path):
        self.path = path
        self.latest: dict[str, dict] = {}
        if path.exists():
            lines = path.read_text(encoding='utf-8').splitlines()
            for n, line in enumerate(lines,1):
                try:
                    item = json.loads(line)
                except json.JSONDecodeError:
                    if n == len(lines):  # 중단된 마지막 줄은 별도 보관 후 제거
                        backup = path.with_name(path.name + '.incomplete-' + datetime.now().strftime('%Y%m%d%H%M%S%f'))
                        backup.write_bytes(path.read_bytes())
                        path.write_text('\n'.join(lines[:-1]) + ('\n' if lines[:-1] else ''), encoding='utf-8')
                        continue
                    raise ValueError(f'실행기록 {n}행이 손상되었습니다.')
                self.latest[item['job_key']] = item

    def reusable(self, job: Job) -> bool:
        item = self.latest.get(job.key, {})
        if item.get('status') != 'saved' or not item.get('artifacts'):
            return False
        if not item.get('confirmed'):
            return False
        for a in item['artifacts']:
            try:
                actual = validate_artifact(Path(a['path']))
            except Exception:
                return False
            if actual['sha256'] != a['sha256']:
                return False
        return True

    def record(self, job: Job, status: str, artifacts: list[dict], confirmed: bool, note: str = '') -> None:
        item = {'timestamp':datetime.now().astimezone().isoformat(), 'job_key':job.key,
                'row':job.client.row, 'service':job.service, 'variant':job.variant,
                'status':status,'confirmed':confirmed,'artifacts':artifacts,'note':note}
        self.path.parent.mkdir(parents=True, exist_ok=True)
        # 미완성 줄이 남아 있을 때 다음 JSON을 같은 줄에 붙이지 않습니다.
        if self.path.exists() and self.path.stat().st_size:
            with self.path.open('rb') as f:
                f.seek(-1,2); tail=f.read(1)
            if tail != b'\n':
                with self.path.open('ab') as f:f.write(b'\n')
        with self.path.open('a',encoding='utf-8') as f:
            f.write(json.dumps(item,ensure_ascii=False)+'\n')
        self.latest[job.key] = item


ROOT = Path(__file__).resolve().parent
URL = 'https://www.hometax.go.kr'
MENUS = {
  1: ('부가가치세 신고도움 서비스','부가가치세 신고도움서비스'),
  2: ('전자(세금)계산서 합계표 조회','전자(세금)계산서 합계표조회'),
  3: ('신용카드/판매(결제)대행 매출자료 조회','신용카드 판매(결제)대행 매출자료 조회'),
  4: ('현금영수증 매출총액 조회',),
  5: ('부가가치세 매입자납부 특례 조회','부가가치세 매입자납부특례 조회'),
  6: ('수출실적명세서 조회',),
  7: ('내국신용장·구매확인서 전자발급명세서조회','내국신용장·구매확인서 전자발급명세서 조회'),
}
FIELD_ALIASES = {
 '사업자등록번호': ('사업자등록번호','사업자번호'),
 '과세연도': ('과세기간(년/기)','과세기간(년)','조회년도','조회연도','귀속연도','결제년도','결제연도','거래년도','선적년월','구매공급일','분기별','일자별'),
 '기수': ('과세기간(년/기)','과세기간','기수','구매공급일'),
 '조회월': ('조회년도','조회연도','월'),
 '시작분기': ('분기','결제년도'), '종료분기': ('분기','결제년도'),
 '시작일': ('조회시작일자','시작일자','시작일','일자별'),
 '종료일': ('조회종료일자','종료일자','종료일','일자별'),
 '조회기간': ('선적년월','조회기간'), '적용분기': ('적용분기',),
 '품목구분': ('품목 구분','품목구분'), '매입매출구분': ('매입/매출 구분','매입매출구분'),
}
RADIO_FIELDS = {'신고구분','분류','구분','작성일자','조회단위','서류구분'}
CONTROL_CSS = 'input:not([type=hidden]):not([type=radio]):not([type=checkbox]),select,[role=combobox],.w2selectbox'


def menu_norm(text):
    return re.sub(r'[\s·ㆍ/()\-–—]', '', str(text))


def choice_norm(text):
    text = norm(text).strip('-')
    if re.fullmatch(r'20\d{2}년?', text):return text.removesuffix('년')
    m = re.fullmatch(r'0*([1-9]\d*)(월|기|분기)', text)
    return str(int(m[1]))+m[2] if m else text


def norm(text):
    return re.sub(r'\s+', '', str(text))


def hometax_host(url):
    host = urlparse(url).hostname or ''
    return host == 'hometax.go.kr' or host.endswith('.hometax.go.kr')


def prompt(message):
    print('\n'+message)
    answer = input('완료 후 Enter / 건너뛰기 s / 종료 q: ').strip().lower()
    while answer not in ('','s','q'):
        answer = input('Enter(완료), s(건너뛰기), q(종료) 중 입력하세요: ').strip().lower()
    if answer == 'q':
        raise KeyboardInterrupt
    if answer == 's':
        raise RuntimeError('사용자가 이 작업을 건너뛰었습니다.')


def chrome_binary(explicit):
    if explicit:
        p = Path(explicit)
        if p.is_file(): return str(p)
        raise ValueError('--chrome 실행파일을 찾지 못했습니다.')
    candidates = [
        Path(os.environ.get('PROGRAMFILES','C:/Program Files'))/'Google/Chrome/Application/chrome.exe',
        Path(os.environ.get('PROGRAMFILES(X86)','C:/Program Files (x86)'))/'Google/Chrome/Application/chrome.exe',
        Path(os.environ.get('LOCALAPPDATA','C:/Users/Default/AppData/Local'))/'Google/Chrome/Application/chrome.exe',
        Path('/Applications/Google Chrome.app/Contents/MacOS/Google Chrome'),
    ]
    for p in candidates:
        if p.is_file(): return str(p)
    for name in ('google-chrome','chromium','chromium-browser'):
        if shutil.which(name): return shutil.which(name)
    raise ValueError('Chrome이 필요합니다. 설치하거나 --chrome으로 실행파일 경로를 지정하세요.')


def debug_info(port):
    # localhost 요청에 조직 HTTP 프록시를 적용하지 않습니다.
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
    try:
        with opener.open(f'http://127.0.0.1:{port}/json/version',timeout=1) as r:
            info = json.load(r)
        return info if 'webSocketDebuggerUrl' in info else None
    except (OSError,ValueError): return None


def launch_chrome(args, incoming):
    profile = args.profile.resolve()
    profile.mkdir(parents=True,exist_ok=True)
    session_file=profile/'vat-browser-session.json'
    existing=debug_info(args.port)
    if existing:
        if session_file.exists():
            session=json.loads(session_file.read_text(encoding='utf-8'))
            return (session.get('websocket')==existing.get('webSocketDebuggerUrl') and
                    session.get('incoming')==str(incoming.resolve()))
        return False
    prefs = profile/'Default'/'Preferences'
    prefs.parent.mkdir(exist_ok=True)
    data = json.loads(prefs.read_text(encoding='utf-8')) if prefs.exists() else {}
    data.setdefault('download',{}).update({'default_directory':str(incoming),'prompt_for_download':False})
    data.setdefault('savefile',{})['default_directory'] = str(incoming)
    sticky = {'recentDestinations':[{'id':'Save as PDF','origin':'local','account':''}],
              'selectedDestinationId':'Save as PDF','version':2,'isHeaderFooterEnabled':False}
    data.setdefault('printing',{}).setdefault('print_preview_sticky_settings',{})['appState']=json.dumps(sticky)
    tmp = prefs.with_suffix('.new')
    tmp.write_text(json.dumps(data),encoding='utf-8'); tmp.replace(prefs)
    subprocess.Popen([chrome_binary(args.chrome),f'--remote-debugging-port={args.port}',
       '--remote-debugging-address=127.0.0.1',f'--user-data-dir={profile}',
       '--no-first-run','--no-default-browser-check','--start-maximized','--kiosk-printing',URL],
       stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    deadline=time.monotonic()+30
    while time.monotonic()<deadline:
        info=debug_info(args.port)
        if info:
            session_file.write_text(json.dumps({'websocket':info['webSocketDebuggerUrl'],
                'incoming':str(incoming.resolve())}),encoding='utf-8')
            return True
        time.sleep(.3)
    raise RuntimeError('Chrome 디버그 연결 시간초과. 전용 프로필을 사용하는 다른 Chrome을 닫으세요.')


class Collector:
    def __init__(self, context, page, args, incoming, kiosk=False):
        self.context,self.page,self.args,self.incoming,self.kiosk = context,page,args,incoming,kiosk
        self.config = json.loads(args.selectors.read_text(encoding='utf-8')) if args.selectors else {}
        self.dialogs=[]
        self.owned=[]
        self.stage='시작'
        self.active_service=None
        self.pending=set()
        self.responses=0
        self.response_errors=0
        context.on('request',self._request)
        context.on('requestfinished',self._request_done)
        context.on('requestfailed',self._request_failed)
        self.context.on('page',self._new_page)
        for p in context.pages: self._watch(p)

    def _request(self,request):
        if request.resource_type in ('xhr','fetch') and hometax_host(request.url):
            self.pending.add(request)

    def _request_done(self,request):
        if request in self.pending:
            self.pending.discard(request)
            response=request.response()
            if response and response.status<400:self.responses+=1
            else:self.response_errors+=1

    def _request_failed(self,request):
        if request in self.pending:
            self.pending.discard(request)
            self.response_errors+=1

    def progress(self,message):
        self.stage=message
        print('  → '+message,flush=True)

    def _watch(self,page):
        def dialog(d):
            # 조회/파일생성 후 안내창. 원문에는 개인정보가 있을 수 있어 로그에 남기지 않습니다.
            self.dialogs.append(d.message)
            d.accept()
        page.on('dialog',dialog)
        page.set_default_timeout(self.args.timeout*1000)

    def _new_page(self,page):
        self._watch(page)
        self.owned.append(page)

    def frames(self,page=None):
        p=page or self.page
        return [f for f in p.frames if f==p.main_frame or f.url in ('','about:blank','about:srcdoc') or hometax_host(f.url)]

    def visible(self,locator):
        return [e for e in locator.all() if e.is_visible()]

    def exact(self,text,page=None,buttons=False,menu=False):
        all_out=[]
        for f in self.frames(page):
            selectors='button,a,label,input[type=button],input[type=submit],[role=button]'
            if not buttons:selectors+=',span,[role=tab]'
            # DOM에서 한 번에 대조하여 큰 메뉴의 요소별 통신 지연을 피합니다.
            ids=f.locator(selectors).evaluate_all(r'''(elements,opts)=>{
                const norm=s=>String(s).replace(opts.menu?/[\s·ㆍ/()\-–—]/g:/\s/g,'');
                return elements.filter(e=>e.getClientRects().length && getComputedStyle(e).visibility!=='hidden'
                    && !(e.tagName==='SPAN' && e.closest('button,a,label,[role=button],[role=tab]'))
                    && norm(e.tagName==='INPUT'?e.value:e.innerText)===norm(opts.text)
                ).map(e=>{if(!e.dataset.vatTarget)e.dataset.vatTarget='vat_'+crypto.randomUUID();return e.dataset.vatTarget;});
            }''',{'text':text,'menu':menu})
            all_out += [f.locator(f'[data-vat-target="{id}"]') for id in ids]
        if len(all_out)==1:return all_out[0]
        if len(all_out)>1:raise RuntimeError(f'같은 이름의 요소 여러 개: {text}. selectors.json으로 범위를 지정하세요.')
        return None

    def wait_click(self,labels,page=None,menu=False,hover=False,timeout=None):
        deadline=time.monotonic()+(timeout if timeout is not None else self.args.timeout)
        while time.monotonic()<deadline:
            for label in labels:
                target=self.exact(label,page,menu=menu)
                if target:
                    target.scroll_into_view_if_needed(timeout=3000)
                    if hover:
                        target.hover(timeout=3000)
                        (page or self.page).wait_for_timeout(350)
                    target.click(timeout=5000)
                    return target
            (page or self.page).wait_for_timeout(200)
        raise RuntimeError('버튼/메뉴 대기시간 초과: '+', '.join(labels))

    def click(self,labels,page=None,required=True):
        for text in labels:
            target=self.exact(text,page,buttons=False)
            if target:
                target.scroll_into_view_if_needed();target.click(timeout=5000)
                return target
        if required:raise RuntimeError('버튼/메뉴를 찾지 못했습니다: '+', '.join(labels))
        return None

    def override(self,job,key,page=None,wait=False,optional=False):
        spec=self.config.get('services',{}).get(str(job.service),{}).get(key)
        if not spec:return None
        end=time.monotonic()+(self.args.timeout if wait else 0)
        while True:
            matches=[]
            for f in self.frames(page):
                if spec.get('frame_url_contains') and spec['frame_url_contains'] not in f.url:continue
                matches+=self.visible(f.locator(spec['css']))
            if matches or not wait or time.monotonic()>=end:break
            (page or self.page).wait_for_timeout(100)
        if 'index' in spec:
            n=int(spec['index'])
            if 0<=n<len(matches):return matches[n]
        if len(matches)==1:return matches[0]
        if not matches and optional:return None
        raise RuntimeError(f'설정 선택자가 한 요소를 지정하지 못했습니다: {key}')

    def field(self,job,key):
        override=self.override(job,'field:'+key)
        if override:return override
        names=tuple(dict.fromkeys((key,*FIELD_ALIASES.get(key,()))))
        matches=[]
        for f in self.frames():
            for name in names:
                matches += self.visible(f.get_by_label(name,exact=True))
            # title/aria-label/placeholder을 정확히 대조하고 전역 검색창은 제외합니다.
            for e in self.visible(f.locator(CONTROL_CSS)):
                labels=[e.get_attribute(a) or '' for a in ('title','aria-label','placeholder')]
                if any(norm(t)==norm(n) for t in labels for n in names):matches.append(e)
            # WebSquare form 표의 라벨 셀 바로 다음 셀만 사용. 문서 전체 following::input 금지.
            for name in names:
                labels=self.visible(f.locator('th,td').filter(has_text=re.compile(r'^\s*'+re.escape(name)+r'\s*$')))
                for label in labels:
                    matches += self.visible(label.locator('xpath=following-sibling::*[1]').locator(CONTROL_CSS))
        # WebSquare의 내부 input과 바깥 div가 같은 컨트롤을 가리키면 바깥 요소로 통합합니다.
        matches=[e.locator('xpath=ancestor-or-self::*[contains(concat(" ",normalize-space(@class)," ")," w2selectbox ")][1]')
                 if e.evaluate('e=>!!e.closest(".w2selectbox")') else e for e in matches]
        unique={e.evaluate("e=>{if(!e.__vat_key)e.__vat_key=crypto.randomUUID();return e.__vat_key}"):e for e in matches}
        matches=list(unique.values())
        # 동일 라벨 아래 연/기수, 연/월, 분기 범위를 값의 형식으로 구분합니다.
        select_matches=[e for e in matches if self.is_combo(e)]
        if key in ('기수','조회월','조회기간','시작분기','종료분기','과세연도','적용분기') and select_matches:
            token={'기수':'기','조회월':'월','조회기간':'기','시작분기':'분기','종료분기':'분기','적용분기':'분기'}.get(key)
            filtered=[]
            for e in select_matches:
                opts=e.locator('option').all_text_contents() if e.evaluate("e=>e.tagName==='SELECT'") else [self.control_value(e)]
                if key=='과세연도' and any(re.fullmatch(r'20\d{2}년?',norm(t)) for t in opts):filtered.append(e)
                elif token and any(token in norm(t) for t in opts):filtered.append(e)
            if filtered:matches=filtered
        if key=='과세연도' and len(matches)>1:
            year_fields=[e for e in matches if re.fullmatch(r'20\d{2}년?',norm(self.control_value(e)))
                         or e.get_attribute('maxlength')=='4']
            if year_fields:matches=year_fields
        if key in ('시작분기','종료분기','시작일','종료일') and len(matches)==2:
            return matches[0 if key.startswith('시작') else 1]
        if len(matches)==1:return matches[0]
        raise RuntimeError(f'{key} 입력란을 정확히 하나 찾지 못했습니다 ({len(matches)}개).')

    def is_combo(self,e):
        return e.evaluate('''e=>e.tagName==='SELECT'||e.getAttribute('role')==='combobox'
            ||e.classList.contains('w2selectbox')||!!e.closest('.w2selectbox')''')

    def control_value(self,e):
        return e.evaluate('''e=>{
            if(e.tagName==='SELECT')return e.options[e.selectedIndex]?.text||'';
            if(e.tagName==='INPUT')return e.value;
            const label=e.querySelector('.w2selectbox_label,input,[role=combobox]');
            return label?(label.value||label.innerText||label.getAttribute('aria-valuetext')||''):
                (e.getAttribute('aria-valuetext')||e.innerText||'');
        }''').strip()

    def select_custom(self,e,value):
        if choice_norm(self.control_value(e))==choice_norm(value):return
        e.click(timeout=5000)
        end=time.monotonic()+min(10,self.args.timeout)
        option=None
        while time.monotonic()<end:
            options=[]
            for frame in self.frames():
                # 펼친 선택지만 대상으로 하여 배경의 라디오 라벨을 누르지 않습니다.
                ids=frame.locator('[role=option],.w2selectbox_item,.w2selectbox_layer td,.w2selectbox_list li,.w2selectbox_options li,[role=listbox] li,[role=listbox] span').evaluate_all(r'''(els,value)=>{
                    const norm=s=>{s=String(s).replace(/\s/g,'').replace(/^-+|-+$/g,'');
                        if(/^20\d{2}년?$/.test(s))return s.replace(/년$/,'');
                        const m=s.match(/^0*([1-9]\d*)(월|기|분기)$/);return m?Number(m[1])+m[2]:s;};
                    let candidates=els.filter(e=>e.getClientRects().length&&getComputedStyle(e).visibility!=='hidden'&&norm(e.innerText)===norm(value));
                    candidates=candidates.filter(e=>!candidates.some(p=>p!==e&&p.contains(e)));
                    return candidates.map(e=>{if(!e.dataset.vatOption)e.dataset.vatOption='vat_'+crypto.randomUUID();return e.dataset.vatOption;});
                }''',value)
                options += [frame.locator(f'[data-vat-option="{id}"]') for id in ids]
            if len(options)>1:raise RuntimeError(f'펼친 선택지에 같은 값이 여러 개입니다: {value}')
            if len(options)==1:option=options[0];break
            self.page.wait_for_timeout(150)
        if not option:
            e.press('Escape')
            raise RuntimeError(f'WebSquare 선택지를 찾지 못했습니다: {value}')
        option.click(timeout=5000)
        end=time.monotonic()+min(5,self.args.timeout)
        while time.monotonic()<end:
            if choice_norm(self.control_value(e))==choice_norm(value):return
            self.page.wait_for_timeout(100)
        raise RuntimeError(f'WebSquare 선택값 확인 실패: {value}')

    def set_field(self,job,key,value):
        if key=='사업자등록번호' and job.service==4 and self.set_split_biz(job,value):return
        if key=='특례조회탭':
            self.click((value,));return
        if key in RADIO_FIELDS:
            target=self.override(job,'field:'+key)
            if target:
                tag=target.evaluate('e=>e.tagName')
                if tag=='SELECT':self.select(target,value);return
                if target.get_attribute('type')=='radio':target.check();return
                target.click();return
            target=self.exact(value)
            if not target:raise RuntimeError(f'{key} 선택을 찾지 못했습니다: {value}')
            target.click()
            linked=target.get_attribute('for')
            if linked:
                checked=target.evaluate('(e,id)=>{const a=e.ownerDocument.getElementById(id);return a? a.checked===true:null}',linked)
                if checked is False:raise RuntimeError(f'{key} 선택 확인 실패')
            return
        e=self.field(job,key)
        tag=e.evaluate('e=>e.tagName')
        if tag=='SELECT':self.select(e,value);return
        if self.is_combo(e):self.select_custom(e,value);return
        if key in ('기수','조회월','적용분기','시작분기','종료분기','조회기간','품목구분','매입매출구분'):
            # 커스텀 콤보박스를 텍스트 input처럼 입력하지 않습니다.
            raise RuntimeError(f'{key}는 커스텀 선택 컨트롤입니다. 선택자 설정/수동확인이 필요합니다.')
        if e.get_attribute('readonly') is not None:raise RuntimeError(f'{key}는 읽기 전용입니다.')
        e.fill(value)
        e.press('Tab')
        actual=e.input_value()
        if re.sub(r'[-\s]','',actual)!=re.sub(r'[-\s]','',value):
            raise RuntimeError(f'{key} 입력값 확인 실패')

    def set_split_biz(self,job,value):
        groups=[]
        for f in self.frames():
            for name in FIELD_ALIASES['사업자등록번호']:
                for label in self.visible(f.locator('th,td').filter(has_text=re.compile(r'^\s*'+re.escape(name)+r'\s*$'))):
                    fields=self.visible(label.locator('xpath=following-sibling::*[1]').locator('input[type=text],input:not([type])'))
                    if len(fields)==3 and [e.get_attribute('maxlength') for e in fields]==['3','2','5']:
                        groups.append(fields)
        if not groups:return False
        if len(groups)!=1:raise RuntimeError('분할 사업자번호 입력란이 여러 개입니다.')
        for e,part in zip(groups[0],(value[:3],value[3:5],value[5:])):
            e.fill(part);e.press('Tab')
            if e.input_value()!=part:raise RuntimeError('분할 사업자번호 입력 확인 실패')
        return True

    def select(self,e,value):
        options=e.locator('option').all()
        exact=[o for o in options if choice_norm(o.inner_text())==choice_norm(value)]
        if len(exact)!=1:raise RuntimeError(f'선택값을 정확히 찾지 못했습니다: {value}')
        v=exact[0].evaluate('e=>e.value')  # HTML option은 value 속성 생략 시 텍스트를 값으로 사용
        e.select_option(value=v)
        if e.input_value()!=v:raise RuntimeError('선택값 확인 실패')

    def wait_busy(self):
        end=time.monotonic()+self.args.timeout
        while time.monotonic()<end:
            busy=self.is_busy()
            if not busy:return
            self.page.wait_for_timeout(200)
        raise RuntimeError('조회 처리중 표시가 사라지지 않았습니다.')

    def is_busy(self):
        # w2modal은 열린 대화상자의 배경에도 쓰이므로 로딩 표시로 간주하지 않습니다.
        return any(f.locator('[aria-busy=true],.w2loading,.w2processbar').evaluate_all(
            "els=>els.some(e=>e.getClientRects().length&&getComputedStyle(e).visibility!=='hidden')") for f in self.frames())

    def search_button(self,job):
        custom=self.override(job,'search')
        if custom:return custom
        labels=('조회',)
        for label in labels:
            e=self.exact(label,buttons=True)
            if e:return e
        raise RuntimeError('해당 조회화면의 조회 버튼을 찾지 못했습니다.')

    def query_screen_ready(self,job):
        if job.service==1:
            try:
                return bool(self.field(job,'사업자등록번호') and self.search_button(job))
            except RuntimeError:return False
        return bool(self.search_button(job))

    def navigate(self,job):
        self.page.bring_to_front()
        self.progress(f'{job.service}번 메뉴 이동: {MENUS[job.service][0]}')
        if self.active_service==job.service and self.query_screen_ready(job):return
        if job.service==1 and self.exact('수임사업자 조회',buttons=True) and self.query_screen_ready(job):
            self.active_service=1;return
        nav=self.override(job,'menu')
        try:
            if nav:nav.click()
            else:
                self.wait_click(('세무대리납세관리',),menu=True,hover=True,timeout=min(self.args.timeout,15))
                self.wait_click(('세무대리 공통조회',),menu=True,hover=True,timeout=min(self.args.timeout,15))
                self.wait_click(MENUS[job.service],menu=True,hover=True,timeout=min(self.args.timeout,15))
            self.wait_busy()
            end=time.monotonic()+self.args.timeout
            while time.monotonic()<end:
                try:
                    if self.query_screen_ready(job):
                        self.active_service=job.service
                        print('  조회화면 진입 완료',flush=True)
                        return
                except RuntimeError:pass
                self.page.wait_for_timeout(200)
            raise RuntimeError('메뉴를 눌렀지만 조회화면 로딩을 확인하지 못했습니다.')
        except Exception:
            if self.args.mode!='assisted':raise
            prompt(f'Chrome에서 [세무대리·납세관리 → 세무대리 공통 조회 → {MENUS[job.service][0]}]를 열어 주세요.')
            self.active_service=job.service

    def select_client(self,job):
        if job.service!=2:return
        # 전자(세금)계산서는 입력란이 아닌 수임사업자 전환 모달을 사용합니다.
        self.click(('수임사업자전환','수임사업자 전환'))
        for f in self.frames():
            candidates=self.visible(f.locator('[role=dialog],.w2popup_window'))
            for modal in candidates:
                if '수임사업자' not in modal.inner_text():continue
                combos=self.visible(modal.locator('select'))
                possible=[]
                for e in combos:
                    if any(norm(t)=='사업자등록번호' for t in e.locator('option').all_text_contents()):possible.append(e)
                if len(possible)!=1:raise RuntimeError('수임사업자번호 검색방식 선택자 확인 필요')
                self.select(possible[0],'사업자등록번호')
                # 번호 입력란은 모달 안으로 한정합니다. ID를 추측하지 않습니다.
                inputs=self.visible(modal.locator('input[type=text],input:not([type])'))
                numbered=[e for e in inputs if '사업자' in (e.get_attribute('title') or '') or e.get_attribute('maxlength') in ('10','12')]
                if len(numbered)!=1:raise RuntimeError('수임사업자 전환 번호 입력란 확인 필요')
                numbered[0].fill(job.client.biz)
                buttons=modal.get_by_text('조회',exact=True)
                if len(self.visible(buttons))!=1:raise RuntimeError('모달 조회 버튼 모호')
                self.visible(buttons)[0].click();self.wait_busy()
                rows=self.visible(modal.locator('tr,[role=row]'))
                rows=[r for r in rows if job.client.biz in re.sub(r'[-\s]','',r.inner_text())]
                if len(rows)!=1:raise RuntimeError('정확한 사업자번호의 수임사업자 결과가 한 건이 아닙니다.')
                radios=self.visible(rows[0].locator('input[type=radio]'))
                if len(radios)!=1:raise RuntimeError('수임사업자 결과 선택 컨트롤 확인 필요')
                radios[0].check()
                confirm=self.visible(modal.get_by_text('확인',exact=True))
                if len(confirm)!=1:raise RuntimeError('수임사업자 확인 버튼 모호')
                confirm[0].click();self.wait_busy();return
        raise RuntimeError('수임사업자 전환 모달을 찾지 못했습니다.')

    def verify_filters(self,job):
        for key,value in job.filters.items():
            if key=='특례조회탭':continue
            if job.service==2 and key in ('사업자등록번호','과세연도'):continue
            if key in RADIO_FIELDS:
                e=self.override(job,'field:'+key) or self.exact(value)
                if not e:raise RuntimeError(f'{key} 선택 상태를 확인하지 못했습니다.')
                state=e.evaluate("""e=>{
                    if(e.tagName==='SELECT')return e.options[e.selectedIndex]?.text;
                    if(e.type==='radio')return e.checked;
                    const input=e.control||e.querySelector('input[type=radio]');
                    if(input)return input.checked;
                    const aria=e.getAttribute('aria-checked')||e.getAttribute('aria-selected');
                    return aria===null?null:aria==='true';
                }""")
                if state is not True and (not isinstance(state,str) or choice_norm(state)!=choice_norm(value)):
                    raise RuntimeError(f'{key} 라디오/탭 선택 상태 확인 실패')
                continue
            if key=='사업자등록번호' and job.service==4:
                # 분할 입력은 입력 직후 이미 대조하며 최종 상태도 확인합니다.
                split=[]
                for frame in self.frames():
                    for cell in self.visible(frame.locator('td')):
                        fields=self.visible(cell.locator('input[type=text],input:not([type])'))
                        if len(fields)==3 and [e.get_attribute('maxlength') for e in fields]==['3','2','5']:
                            split.append(''.join(e.input_value() for e in fields))
                if split:
                    if split!=[value]:raise RuntimeError('조회 전 분할 사업자번호가 변경되었습니다.')
                    continue
            e=self.field(job,key)
            actual=self.control_value(e)
            if self.is_combo(e):same=choice_norm(actual)==choice_norm(value)
            else:same=re.sub(r'[-.\s]','',actual)==re.sub(r'[-.\s]','',value)
            if not same:raise RuntimeError(f'조회 전 {key} 조건이 변경되었습니다.')

    def body_signature(self):
        # input 값은 innerText에 포함되지 않아 입력값만 바뀐 것을 결과 변화로 오인하지 않습니다.
        return '\n'.join(f.locator('body').inner_text() for f in self.frames())

    def query_error(self):
        return any(any(t in message for t in ('권한','로그인','유효하지','잘못','필수','실패','오류'))
                   for message in self.dialogs)

    def wait_query(self,job,previous,baseline,error_baseline):
        end=time.monotonic()+self.args.timeout
        quiet_since=None
        last_status=time.monotonic()
        while time.monotonic()<end:
            if self.response_errors>error_baseline:
                raise RuntimeError('홈택스 조회 요청이 실패했습니다. 연결/응답을 확인하세요.')
            if self.query_error():raise RuntimeError('홈택스가 조회 오류/권한/로그인 알림을 표시했습니다. 화면을 확인하세요.')
            busy=self.is_busy()
            changed=self.body_signature()!=previous
            finished=self.responses>baseline
            output=any(self.exact(label,buttons=True) for label in
                ('인쇄하기','인쇄','명세서 출력','출력','파일생성','파일 생성','엑셀 내려받기','엑셀내려받기'))
            empty=any('없' in m and ('자료' in m or '결과' in m) for m in self.dialogs) or self.no_data(job)
            if not busy and not self.pending and (changed or finished or empty) and (output or empty):
                if quiet_since is None:quiet_since=time.monotonic()
                if time.monotonic()-quiet_since>=1:
                    return
            else:quiet_since=None
            if time.monotonic()-last_status>=10:
                print('  조회 응답/결과를 기다리는 중… (Enter 입력 불필요)',flush=True)
                last_status=time.monotonic()
            self.page.wait_for_timeout(200)
        raise RuntimeError('조회 완료를 확인하지 못했습니다. 응답/결과 변화와 출력 버튼 확인 시간초과')

    def configure_query(self,job):
        self.progress(f'번호 끝 {job.client.biz[-4:]} / {job.client.period.label}: 조회조건 설정')
        failures=[]
        try:self.select_client(job)
        except Exception as exc:
            if self.args.mode!='assisted':raise
            failures.append(str(exc))
        ordered=list(job.filters.items())
        ordered.sort(key=lambda item: 0 if item[0]=='특례조회탭' else (1 if item[0] in RADIO_FIELDS else 2))
        for key,value in ordered:
            if job.service==2 and key in ('사업자등록번호','과세연도'):continue
            try:
                self.set_field(job,key,value)
                self.wait_busy()
            except Exception as exc:
                if self.args.mode!='assisted':raise
                failures.append(str(exc))
        if failures:
            print('자동 설정하지 못한 항목:', '; '.join(failures),flush=True)
            masked={k:('끝 4자리 '+v[-4:] if k=='사업자등록번호' else v) for k,v in job.filters.items()}
            print('필수 조건:',json.dumps(masked,ensure_ascii=False),flush=True)
            prompt('사업자번호와 조건을 모두 설정하세요. 수임사업자 전환이 필요하면 완료하세요.')
        if self.args.mode!='assisted':self.verify_filters(job)
        self.progress('조회 버튼 클릭 및 처리 완료 대기')
        self.dialogs.clear()
        previous=self.body_signature()
        baseline=self.responses
        error_baseline=self.response_errors
        self.search_button(job).click()
        ready=self.override(job,'result_ready',wait=True)
        if ready:
            ready.wait_for(state='visible')
            self.wait_busy()
        elif self.args.mode=='auto':
            self.wait_query(job,previous,baseline,error_baseline)
        else:
            self.page.wait_for_timeout(600)
            self.wait_busy()
        if self.query_error():raise RuntimeError('홈택스가 조회 오류/권한/로그인 알림을 표시했습니다. 화면을 확인하세요.')
        if self.args.mode=='assisted':
            prompt(f'[{job.client.company} / 번호 끝 {job.client.biz[-4:]} / {job.client.period.label} / {job.variant}]\n조회결과의 거래처·기간 및 무자료 여부를 확인하세요.')
            return True
        proof=self.override(job,'result_identity')
        if self.args.mode=='strict' and not ready:
            raise RuntimeError('strict 모드는 result_ready와 result_identity 선택자가 필요합니다.')
        if proof:
            self.verify_identity(job,proof)
            return True
        if self.args.mode=='strict':
            raise RuntimeError('strict 모드는 result_ready와 result_identity 선택자가 필요합니다.')
        # 자동 수집은 계속하되 입력값만으로 결과 내용까지 확인되었다고 기록하지 않습니다.
        return False

    def verify_identity(self,job,element):
        text=re.sub(r'[-\s]','',element.inner_text())
        if job.client.biz not in text or str(job.client.period.year) not in text:
            raise RuntimeError('결과 영역에서 사업자번호/연도를 확인하지 못했습니다.')
        for key,value in job.filters.items():
            if key in ('사업자등록번호','과세연도','작성일자','특례조회탭'):continue
            if re.sub(r'[-\s]','',value) not in text:
                raise RuntimeError(f'결과 영역에서 {key} 조건을 확인하지 못했습니다.')

    def no_data(self,job):
        custom=self.override(job,'no_data',optional=True)
        if custom:return True
        # 3번은 한 표가 비어도 다른 표에 매출이 있으므로 통합 무자료로 판단하지 않습니다.
        if job.service==3:return False
        if any('없' in m and ('조회된 자료' in m or '조회 결과' in m or '조회결과' in m) for m in self.dialogs):return True
        bodies=[f.locator('body').inner_text() for f in self.frames()]
        text='\n'.join(bodies)
        total=re.findall(r'총\s*([0-9,]+)\s*건',text)
        if total:return all(int(n.replace(',',''))==0 for n in total)
        # 모호한 공통 '조회결과 없음'은 판정하지 않고 일반 출력 절차를 거칩니다.
        return False

    def download(self,job,target,labels):
        button=self.override(job,'download')
        if not button:
            for label in labels:
                button=self.exact(label)
                if button:break
        if not button:raise RuntimeError('다운로드 버튼을 찾지 못했습니다.')
        with self.page.expect_download(timeout=self.args.timeout*1000) as event:
            button.click()
        download=event.value
        if download.failure():raise RuntimeError('브라우저 다운로드 실패')
        name=Path(download.suggested_filename).name
        suffix=Path(name).suffix.lower()
        if not suffix:raise RuntimeError('다운로드 파일의 확장자가 없습니다.')
        dest=target/(job.stem+suffix)
        if dest.exists():raise RuntimeError('같은 이름의 새 다운로드가 중복됩니다.')
        download.save_as(str(dest))
        return self.artifact(dest,'native_download')

    def generate_download(self,job,target):
        generation=self.override(job,'generate')
        if generation:generation.click()
        else:self.click(('파일생성','파일 생성'))
        # 파일 생성은 화면 안내대로 시간이 걸립니다. 재클릭하지 않고 활성화/안내를 기다립니다.
        end=time.monotonic()+self.args.generation_timeout
        saw_disabled=False
        while time.monotonic()<end:
            if any(any(t in m for t in ('실패','오류','권한')) for m in self.dialogs):
                raise RuntimeError('파일생성 오류. 홈택스 알림을 확인하세요.')
            button=self.override(job,'download',optional=True)
            if not button:button=self.exact('다운로드')
            enabled=bool(button and button.is_enabled() and button.get_attribute('aria-disabled')!='true')
            if not enabled:saw_disabled=True
            completed=any('파일생성' in norm(m) and '완료' in m for m in self.dialogs)
            if enabled and (saw_disabled or completed):
                return self.download(job,target,('다운로드',))
            self.page.wait_for_timeout(500)
        raise RuntimeError('파일생성 시간초과. 다음 실행에서 다시 조회/생성할 수 있습니다.')

    def artifact(self,path,source):
        return {'path':str(path.resolve()),'source':source,**validate_artifact(path)}

    def screen_pdf(self,job,target,no_data=False):
        # 보고서 PDF와 구별되는 조회화면 증빙입니다. 테이블의 전체 페이지를 대체하지 않습니다.
        session=self.context.new_cdp_session(self.page)
        try:
            import base64
            result=session.send('Page.printToPDF',{'printBackground':True,'landscape':True,'preferCSSPageSize':False})
            path=target/(job.stem+'_조회화면증빙.pdf')
            path.write_bytes(base64.b64decode(result['data']))
        finally:session.detach()
        return self.artifact(path,'no_data_screen' if no_data else 'query_screen')

    def report_pdf(self,job,target):
        self.progress('전체 보고서 열기 및 PDF 저장')
        if self.args.mode!='assisted' and not self.kiosk:
            raise RuntimeError('자동 PDF 설정이 확인되지 않은 기존 Chrome입니다. 홈택스 전용 Chrome을 닫고 다시 실행하세요.')
        started_at=time.time()
        incoming_before={p for p in self.incoming.iterdir() if p.is_file()}
        pages_before=set(self.context.pages)
        button=self.override(job,'report')
        if button:button.click()
        elif job.service==2:self.click(('명세서 출력','출력'))
        else:self.click(('인쇄하기','인쇄'))
        end=time.monotonic()+10
        while time.monotonic()<end and not (set(self.context.pages)-pages_before):
            self.page.wait_for_timeout(200)
        new=list(set(self.context.pages)-pages_before)
        if len(new)>1:raise RuntimeError('여러 리포트창이 열렸습니다. 출력 대상을 확인하세요.')
        report=new[0] if new else self.page
        report.bring_to_front()
        report.wait_for_load_state('domcontentloaded')
        # WebSquare/CLIP report 로딩 완료 확인. 제목만 열린 빈 팝업은 성공 처리하지 않습니다.
        report.wait_for_timeout(1500)
        if self.args.mode=='strict':
            ready=self.override(job,'report_ready',report,wait=True)
            identity=self.override(job,'report_identity',report)
            if not ready or not identity:raise RuntimeError('strict 모드는 report_ready와 report_identity 선택자가 필요합니다.')
            self.verify_identity(job,identity)
            if not self.kiosk:raise RuntimeError('strict PDF 인쇄는 이 실행에서 준비한 전용 Chrome이 필요합니다.')
        if not hometax_host(report.url) and report.url not in ('','about:blank'):
            raise RuntimeError('홈택스 외부 리포트창입니다. URL을 확인하세요.')
        if self.args.mode=='assisted':
            prompt('리포트의 거래처·기간·전체 페이지가 맞고 로딩이 끝났는지 확인하세요. 아직 인쇄하지 말고 Enter를 누르세요.')
        print_button=self.override(job,'print',report)
        if not print_button:
            # 프린터 아이콘의 접근성 속성만 사용합니다. 임의 좌표 클릭 금지.
            deadline=time.monotonic()+self.args.timeout
            while time.monotonic()<deadline:
                print_button=self.exact('인쇄',report,buttons=True)
                if print_button:break
                icons=[]
                for f in self.frames(report):
                    icons+=self.visible(f.locator('button[title*="인쇄"],a[title*="인쇄"],button[aria-label*="인쇄"],a[aria-label*="인쇄"],button[class*="print"],img[alt*="인쇄"]:not([hidden])'))
                if len(icons)>1:raise RuntimeError('리포트 프린터 아이콘이 여러 개입니다. print 선택자를 확인하세요.')
                if len(icons)==1:
                    icons[0].click(timeout=5000)
                    report.wait_for_timeout(500)
                    print_button=self.exact('인쇄',report,buttons=True)
                    if print_button:break
                report.wait_for_timeout(200)
        if print_button:
            # 전체페이지 라디오가 존재하면 명확히 선택합니다.
            full=self.exact('전체 페이지',report) or self.exact('전체페이지',report)
            if full:full.click()
            if not self.kiosk and self.args.mode=='assisted':
                print('이 Chrome은 연결된 기존 창입니다. 인쇄 미리보기에서 [PDF로 저장]을 선택하세요.')
            print_button.click()
            if not self.kiosk and self.args.mode=='assisted':
                prompt(f'인쇄 미리보기에서 [PDF로 저장]으로 저장하세요. 저장 위치: {self.incoming}. 저장 완료 후 Enter.')
        else:
            if self.args.mode!='assisted':raise RuntimeError('리포트 인쇄 버튼을 찾지 못했습니다. print 선택자 확인 필요')
            prompt('리포트 프린터 아이콘 → 전체페이지 → 인쇄를 눌러 주세요. Chrome에서는 [PDF로 저장]을 선택하세요.\n새 PDF는 실행 폴더의 _incoming에 저장하세요. PDF 준비 후 Enter.')
        try:
            return self.wait_print_file(job,target,incoming_before)
        except Exception:
            if self.args.mode!='assisted':raise
            print('자동 PDF 저장이 확인되지 않았습니다.')
            path=input('PDF를 수동 저장한 다음 파일의 전체 경로 입력 (s: 건너뛰기): ').strip().strip('"')
            if path.lower()=='s':raise RuntimeError('리포트 저장 건너뜀')
            source=Path(path)
            validate_artifact(source)
            if source.suffix.lower()!='.pdf':raise RuntimeError('PDF만 지정할 수 있습니다.')
            if source.resolve() in {p.resolve() for p in incoming_before} or source.stat().st_mtime < started_at-2:
                raise RuntimeError('이전 작업 파일은 사용할 수 없습니다. 이번 조회 PDF를 새로 저장하세요.')
            dest=target/(job.stem+'.pdf');shutil.copy2(source,dest)
            return self.artifact(dest,'manual_report')
        finally:
            for p in new:
                if not p.is_closed():p.close()
            if not self.page.is_closed():self.page.bring_to_front()

    def wait_print_file(self,job,target,before):
        end=time.monotonic()+self.args.timeout
        previous={}
        while time.monotonic()<end:
            candidates=[p for p in self.incoming.iterdir() if p.is_file() and p not in before and p.suffix.lower()=='.pdf']
            if len(candidates)>1:raise RuntimeError('새 PDF가 여러 개여서 출력 대상을 판정할 수 없습니다.')
            if len(candidates)==1:
                p=candidates[0];size=p.stat().st_size
                if size>0 and previous.get(p)==size:
                    validate_artifact(p)
                    dest=target/(job.stem+'.pdf');shutil.move(str(p),str(dest))
                    return self.artifact(dest,'clipreport_print')
                previous[p]=size
            self.page.wait_for_timeout(500)
        raise RuntimeError('PDF 저장 시간초과')

    def collect(self,job,target):
        self.navigate(job)
        confirmed=self.configure_query(job)
        artifacts=[]
        empty=self.no_data(job)
        notes=[]
        if not confirmed:
            notes.append('조회조건 선택값과 응답 완료 확인; 보고서 내용의 사업자·기간 대조는 미확인')
        # 특례 탭 외에는 화면 전체를 원본 데이터인 것처럼 XLSX로 변환하지 않습니다.
        if '원본' in job.client.formats and job.service in (5,6,7):
            if empty:notes.append('조회 건수 0: 원본 다운로드를 시도하지 않았음')
            elif job.service in (6,7):artifacts.append(self.generate_download(job,target))
            else:artifacts.append(self.download(job,target,('엑셀 내려받기','엑셀내려받기','엑셀 다운로드','엑셀다운로드')))
        if 'PDF' in job.client.formats:
            if job.service==5 or empty:
                artifacts.append(self.screen_pdf(job,target,no_data=empty))
                notes.append('조회화면 증빙 PDF: 전체 거래 원본/공식 보고서를 대체하지 않음')
            else:artifacts.append(self.report_pdf(job,target))
        if '원본' in job.client.formats and job.service in (1,2,3,4):
            notes.append('이 구현은 해당 항목의 PDF를 수집함; 엑셀 원본 미수집')
        if not artifacts:raise RuntimeError('요청 형식에 대해 수집한 파일이 없습니다. PDF 또는 모두로 지정하세요.')
        return artifacts,confirmed,'; '.join(notes)

    def save_diagnostics(self,job,target,exc):
        # 값/텍스트/비밀번호/쿠키/스크린샷 없이 컨트롤 구조만 남깁니다.
        record={'stage':self.stage,'service':job.service,'row':job.client.row,
                'error_type':type(exc).__name__,'frames':[]}
        if isinstance(exc,(RuntimeError,ValueError)):record['reason']=str(exc)
        try:
            for frame in self.frames():
                controls=frame.locator('input:not([type=password]):not([type=hidden]),select,button,a,.w2selectbox,[role=combobox]').evaluate_all('''els=>els.filter(e=>e.getClientRects().length).map(e=>({
                    tag:e.tagName,id:e.id,class:e.className,type:e.type||'',role:e.getAttribute('role'),
                    title:e.getAttribute('title'),aria_label:e.getAttribute('aria-label'),maxlength:e.getAttribute('maxlength'),
                    disabled:e.disabled||false,readonly:e.readOnly||false,
                    options:e.tagName==='SELECT'?Array.from(e.options).map(o=>o.text).filter(s=>/^(20\\d{2}년?|[12]기|0?\\d{1,2}월|[1-4]분기)$/.test(s.trim())):undefined
                }))''')
                record['frames'].append({'path':urlparse(frame.url).path,'controls':controls})
        except Exception:record['capture_error']='제어요소 구조를 읽지 못함'
        path=target/'오류진단.json'
        path.write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
        return path


def parser():
    p=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--input',type=Path,help='기본: 코드 폴더의 확장 엑셀 또는 원래 INPUT-부가세.xlsx')
    p.add_argument('--sheet')
    p.add_argument('--year',type=int,default=2026,help='연도 열이 없을 때 적용. 기본 2026. 엑셀 값이 우선.')
    p.add_argument('--output',type=Path,default=ROOT/'다운로드')
    p.add_argument('--profile',type=Path,default=ROOT/'.chrome-profile')
    p.add_argument('--chrome',help='Chrome 실행파일 경로')
    p.add_argument('--port',type=int,default=9222)
    p.add_argument('--mode',choices=('auto','assisted','strict'),default='auto')
    p.add_argument('--selectors',type=Path,help='실제 화면에서 확인한 선택자 JSON')
    p.add_argument('--timeout',type=int,default=60)
    p.add_argument('--generation-timeout',type=int,default=360)
    p.add_argument('--validate-only',action='store_true')
    p.add_argument('--services',help='이번 실행의 항목만 선택, 예: 1,2,7')
    p.add_argument('--inspect',action='store_true',help='로그인 후 현재 화면의 제어요소 메타데이터 저장. 입력값/쿠키 제외.')
    p.add_argument('--resume',action='store_true',help='저장기록·확인상태·파일 해시가 유효한 작업만 건너뜀')
    p.add_argument('--continue-on-error',action='store_true',help='오류 기록 후 다음 작업 계속. 기본은 첫 오류에서 멈춤.')
    return p


def discover_input(explicit=None):
    if explicit:
        if not explicit.is_file():raise ValueError(f'입력 엑셀 파일을 찾지 못했습니다: {explicit}')
        return explicit.resolve()
    for name in ('INPUT-부가세_확장.xlsx','INPUT-부가세.xlsx','INPUT_VAT_2026.xlsx'):
        if (ROOT/name).is_file():return ROOT/name
    candidates=[p for p in ROOT.glob('*.xlsx') if not p.name.startswith('~$') and '빈양식' not in p.name]
    if len(candidates)==1:return candidates[0]
    raise ValueError('코드와 같은 폴더에 INPUT-부가세.xlsx를 두세요. 다른 파일명은 --input "파일명.xlsx"로 지정하세요.')


def main(argv=None):
    # CMD에서도 진행 메시지를 즉시 표시합니다.
    for stream in (sys.stdout,sys.stderr):
        if hasattr(stream,'reconfigure'):stream.reconfigure(line_buffering=True)
    args=parser().parse_args(argv)
    if not 1024<=args.port<=65535 or min(args.timeout,args.generation_timeout)<=0:
        print('포트/시간 제한값이 잘못되었습니다.',file=sys.stderr);return 2
    try:
        args.input=discover_input(args.input)
        print(f'입력파일: {args.input.name} / 기본 보완 연도: {args.year} / 모드: {args.mode}',flush=True)
        clients=load_clients(args.input,args.year,args.sheet)
        if args.services:
            ids={int(x) for x in args.services.split(',')}
            if not ids or not ids.issubset(SERVICES):raise ValueError('--services는 1~7이어야 합니다.')
            clients=[replace(c,services=tuple(s for s in c.services if s in ids)) for c in clients]
            clients=[c for c in clients if c.services]
        jobs=build_jobs(clients)
        if not jobs:raise ValueError('선택된 작업이 없습니다.')
        if any(j.client.formats==('원본',) and j.service in (1,2,3,4) for j in jobs):
            raise ValueError('1~4번은 이 구현에서 원본 전용 저장을 지원하지 않습니다. 저장형식을 PDF 또는 모두로 지정하세요.')
        print(f'거래처 {len(clients)}곳, 세부 작업 {len(jobs)}건')
        for j in jobs:
            print(f'  엑셀 {j.client.row}행 / {j.client.period.label} / {j.service} / {j.variant}')
        if args.validate_only:
            print('입력·조회계획 검증 완료. 홈택스 로그인/조회/다운로드는 실행하지 않았습니다.');return 0
        from playwright.sync_api import sync_playwright
        args.output=args.output.resolve();args.output.mkdir(parents=True,exist_ok=True)
        incoming=args.output/'_incoming';incoming.mkdir(exist_ok=True)
        journal=Journal(args.output/'실행기록.jsonl')
        success=failed=skipped=unverified=0
        with sync_playwright() as pw:
            kiosk=launch_chrome(args,incoming)
            if args.mode!='assisted' and not kiosk:
                raise RuntimeError('이전 버전의 전용 Chrome이 열려 있습니다. 홈택스 전용 Chrome 창을 닫은 뒤 다시 실행하세요. 일반 Chrome 창은 그대로 두어도 됩니다.')
            browser=pw.chromium.connect_over_cdp(f'http://127.0.0.1:{args.port}')
            context=browser.contexts[0]
            pages=[p for p in context.pages if hometax_host(p.url) and 'clipreport' not in p.url]
            page=pages[0] if pages else context.new_page()
            if not pages:page.goto(URL,wait_until='domcontentloaded')
            print('전용 Chrome에서 [세무대리인] 로그인 및 필요한 수임 동의를 완료하세요. 로그인되면 메뉴 이동부터 자동 진행합니다. Enter 입력은 필요 없습니다.')
            end=time.monotonic()+600
            last_status=time.monotonic()
            while time.monotonic()<end:
                logged=[]
                for p in context.pages:
                    if p.is_closed() or not hometax_host(p.url) or 'clipreport' in p.url:continue
                    if any(f.get_by_text('로그아웃',exact=True).count() and any(e.is_visible() for e in f.get_by_text('로그아웃',exact=True).all()) for f in p.frames):logged.append(p)
                if logged:
                    page=logged[-1]
                    print('로그인 확인 완료. 자동 처리를 시작합니다.',flush=True)
                    break
                if time.monotonic()-last_status>=20:
                    print('로그인을 기다리는 중… (Chrome에서 로그인 / 중단은 Ctrl+C)',flush=True)
                    last_status=time.monotonic()
                page.wait_for_timeout(1000)
            else:raise RuntimeError('로그인 대기 10분 초과')
            collector=Collector(context,page,args,incoming,kiosk)
            if args.inspect:
                prompt('선택자를 확인할 조회화면을 열어 주세요. 고객 자료를 입력할 필요는 없습니다.')
                controls=[]
                for frame in collector.frames():
                    for e in collector.visible(frame.locator('input,select,button,a,[role=combobox]')):
                        if e.get_attribute('type') in ('password','hidden'):continue
                        item={k:e.get_attribute(k) for k in ('id','name','title','aria-label','role','type','maxlength')}
                        item['tag']=e.evaluate('e=>e.tagName')
                        item['frame_path']=urlparse(frame.url).path
                        if item['tag']=='SELECT':item['options']=e.locator('option').all_text_contents()
                        # 버튼·링크 텍스트에는 고객명이 있을 수 있어 기본 수집하지 않습니다.
                        controls.append(item)
                dest=args.output/('controls_'+time.strftime('%Y%m%d_%H%M%S')+'.json')
                dest.write_text(json.dumps(controls,ensure_ascii=False,indent=2),encoding='utf-8')
                print(f'선택자 점검자료 저장: {dest}. 공유 전 개인정보 포함 여부를 확인하세요.')
                return 0
            for job in jobs:
                if args.resume and journal.reusable(job):
                    skipped+=1;print(f'재사용: {job.service}/{job.variant}');continue
                # 고유 실행 폴더: 기존 파일을 삭제/덮어쓰지 않습니다.
                target=args.output/job.key/time.strftime('%Y%m%d_%H%M%S')
                n=0
                while target.exists():
                    n+=1;target=args.output/job.key/(time.strftime('%Y%m%d_%H%M%S')+f'_{n}')
                target.mkdir(parents=True)
                metadata={'job_key':job.key,'row':job.client.row,'company':job.client.company,
                          'service':job.service,'variant':job.variant,'filters':job.filters,
                          'formats':job.client.formats}
                (target/'조회조건.json').write_text(json.dumps(metadata,ensure_ascii=False,indent=2),encoding='utf-8')
                print(f'처리: 엑셀 {job.client.row}행 / 항목 {job.service} / {job.variant}')
                try:
                    artifacts,confirmed,note=collector.collect(job,target)
                    journal.record(job,'saved',artifacts,confirmed,note);success+=1
                    if not confirmed:unverified+=1
                    print(f'저장 완료: {len(artifacts)}개 / {target}',flush=True)
                    if not confirmed:print('  보고서 내용의 사업자·기간은 확인 필요. 자동 검증 성공으로 기록하지 않았습니다.',flush=True)
                except KeyboardInterrupt:
                    journal.record(job,'interrupted',[],False,'사용자가 중단함')
                    raise
                except Exception as exc:
                    # 브라우저 예외 본문에 DOM/개인정보가 포함될 수 있어 원문 로그를 저장하지 않습니다.
                    journal.record(job,'failed',[],False,type(exc).__name__)
                    failed+=1
                    detail=str(exc) if isinstance(exc,(RuntimeError,ValueError)) else '브라우저 요소/대기 오류. 해당 화면과 선택자를 확인하세요.'
                    print(f'실패: {type(exc).__name__}: {detail}')
                    diagnostic=collector.save_diagnostics(job,target,exc)
                    print(f'오류진단: {diagnostic}',flush=True)
                    if not args.continue_on_error:
                        print('첫 오류에서 멈춥니다. 위 원인과 오류진단.json을 확인하세요. 입력 대기 상태가 아닙니다.',flush=True)
                        break
                    print('다음 작업으로 넘어갑니다. 실패한 작업은 --resume 실행 시 다시 처리합니다.')
                finally:
                    for popup in list(collector.owned):
                        if popup!=page and not popup.is_closed():popup.close()
                    collector.owned.clear()
                    if page.is_closed():raise RuntimeError('메인 홈택스 창이 닫혔습니다.')
                    page.bring_to_front()
            # 로그인된 사용자 브라우저는 그대로 두고 Playwright 연결만 문맥 종료로 해제합니다.
        print(f'완료: 저장 {success}(내용 확인 필요 {unverified}), 실패 {failed}, 이전 파일 재사용 {skipped}')
        print(f'결과/실행기록: {args.output}')
        return 1 if failed else 0
    except KeyboardInterrupt:
        print('\n중단되었습니다. 다음 실행에 --resume을 사용하세요.');return 130
    except Exception as exc:
        # 입력/초기화 오류는 비밀값을 받지 않으므로 원인을 출력합니다.
        print(f'실행 중단: {exc}',file=sys.stderr);return 2


if __name__=='__main__':
    raise SystemExit(main())
