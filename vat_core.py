"""홈택스 입력/업무계획/산출물 검증. 브라우저 없이 사용할 수 있습니다."""
from __future__ import annotations
import calendar
import hashlib
import json
import re
from dataclasses import dataclass, asdict
from datetime import date, datetime
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

SERVICES = {
    1: '부가가치세 신고도움서비스',
    2: '전자(세금)계산서 합계표조회',
    3: '신용카드 판매(결제)대행 매출자료 조회',
    4: '현금영수증 매출총액 조회',
    5: '부가가치세 매입자납부 특례 조회',
    6: '수출실적명세서 조회',
    7: '내국신용장 구매확인서 전자발급명세서 조회',
}
REQUIRED = ('사업자명', '사업자등록번호', '과세기간')


def clean(value: Any) -> str:
    return '' if value is None else str(value).strip()


def choices(value: Any, allowed: tuple[str, ...], default: str) -> tuple[str, ...]:
    text = clean(value) or default
    if text == '모두':
        return allowed
    parts = tuple(x.strip() for x in re.split('[,;|]', text) if x.strip())
    if not parts or len(parts) != len(set(parts)) or any(p not in allowed for p in parts):
        raise ValueError(f'허용값 {allowed} 또는 모두: {text}')
    return parts


def biz_number(value: Any) -> str:
    # 숫자로 저장된 앞자리 0은 원본 소실 여부를 판단할 수 없으므로 추측하지 않습니다.
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r'[-\s]', '', clean(value))
    if not re.fullmatch(r'[0-9]{10}', text):
        raise ValueError('사업자등록번호는 10자리 숫자(하이픈 허용)여야 합니다. 셀을 텍스트로 저장하세요.')
    return text


@dataclass(frozen=True)
class Period:
    year: int
    term: int
    filing: str
    scope: str

    @classmethod
    def parse(cls, text: Any, year: Any, scope: Any = None) -> 'Period':
        compact = re.sub(r'\s+', '', clean(text))
        m = re.fullmatch(r'([12])기(예정|확정)', compact)
        if not m:
            raise ValueError('과세기간은 1기예정/1기확정/2기예정/2기확정 중 하나여야 합니다.')
        if not re.fullmatch(r'[0-9]{4}', clean(year)):
            raise ValueError('과세연도 4자리가 필요합니다. 엑셀에 추가하거나 --year로 지정하세요.')
        y = int(clean(year))
        if not 2000 <= y <= 2100:
            raise ValueError('과세연도는 2000~2100 범위여야 합니다.')
        s = clean(scope) or ('분기' if m[2] == '예정' else '반기')
        if s not in ('분기', '반기'):
            raise ValueError('조회범위는 분기 또는 반기여야 합니다.')
        if m[2] == '예정' and s != '분기':
            raise ValueError('예정신고 조회범위는 분기여야 합니다.')
        return cls(y, int(m[1]), m[2], s)

    @property
    def quarters(self) -> tuple[int, ...]:
        base = 1 if self.term == 1 else 3
        if self.filing == '예정':
            return (base,)
        return (base, base + 1) if self.scope == '반기' else (base + 1,)

    @property
    def months(self) -> tuple[int, ...]:
        return tuple(m for q in self.quarters for m in range(q * 3 - 2, q * 3 + 1))

    @property
    def start(self) -> str:
        return date(self.year, self.months[0], 1).isoformat()

    @property
    def end(self) -> str:
        last = self.months[-1]
        return date(self.year, last, calendar.monthrange(self.year, last)[1]).isoformat()

    @property
    def label(self) -> str:
        return f'{self.year}_{self.term}기{self.filing}_{self.scope}'


@dataclass(frozen=True)
class Client:
    row: int
    company: str
    biz: str
    period: Period
    services: tuple[int, ...]
    invoice_types: tuple[str, ...]
    directions: tuple[str, ...]
    documents: tuple[str, ...]
    special_tabs: tuple[str, ...]
    goods: str
    special_direction: str
    formats: tuple[str, ...]


@dataclass(frozen=True)
class Job:
    client: Client
    service: int
    variant: str
    filters: dict[str, str]

    @property
    def key(self) -> str:
        raw = json.dumps({'biz':self.client.biz, 'period':asdict(self.client.period),
                          'formats':self.client.formats, 'service':self.service,
                          'variant':self.variant, 'filters':self.filters}, sort_keys=True, ensure_ascii=False)
        return hashlib.sha256(raw.encode()).hexdigest()[:24]

    @property
    def stem(self) -> str:
        # 전체 사업자번호를 파일명/일반 로그에 노출하지 않습니다.
        raw = f'{self.client.company}_{self.client.biz[-4:]}_{self.client.period.label}_{self.service:02d}_{SERVICES[self.service]}_{self.variant}'
        name = re.sub(r'[\\/:*?"<>|\x00-\x1f]', '_', raw).strip(' .')
        return name[:150] + '_' + self.key[:8]


def load_clients(path: Path, fallback_year: int | None = None, sheet: str | None = None) -> list[Client]:
    w = load_workbook(path, data_only=False, read_only=True)
    try:
        s = w[sheet] if sheet else w.active
        rows = iter(s.iter_rows())
        headers = [clean(c.value) for c in next(rows)]
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise ValueError('엑셀 헤더가 중복됩니다.')
        for key in REQUIRED:
            if key not in headers:
                raise ValueError(f'필수 열 없음: {key}')
        clients: list[Client] = []
        seen = set()
        for rownum, row in enumerate(rows, 2):
            if not any(c.value is not None for c in row):
                continue
            if any(c.data_type == 'f' for c in row):
                raise ValueError(f'{rownum}행: 수식 대신 확정된 입력값을 사용하세요.')
            data = dict(zip(headers, (c.value for c in row)))
            if clean(data.get('실행여부')).upper() in ('N', '아니오', '0'):
                continue
            try:
                company = clean(data.get('사업자명'))
                if not company:
                    raise ValueError('사업자명 누락')
                p = Period.parse(data.get('과세기간'), data.get('과세연도') or fallback_year, data.get('조회범위'))
                svc = clean(data.get('조회항목')) or '1,2,3,4,5,6,7'
                services = tuple(int(x.strip()) for x in re.split('[,;|]', svc))
                if not services or len(services) != len(set(services)) or any(x not in SERVICES for x in services):
                    raise ValueError('조회항목은 1~7 숫자를 쉼표로 구분하세요.')
                goods = clean(data.get('특례품목')) or '전체'
                sd = clean(data.get('특례매입매출구분')) or '전체'
                if sd not in ('전체', '매입', '매출'):
                    raise ValueError('특례매입매출구분은 전체/매입/매출')
                client = Client(rownum, company, biz_number(data.get('사업자등록번호')), p, services,
                    choices(data.get('계산서종류'), ('전자세금계산서','전자계산서'), '모두'),
                    choices(data.get('매출매입구분'), ('매출','매입'), '모두'),
                    choices(data.get('무역서류종류'), ('내국신용장','구매확인서'), '모두'),
                    choices(data.get('특례조회탭'), ('거래내역조회','전용계좌조회','국고입금내역','기납부세액'), '거래내역조회'),
                    goods, sd, choices(data.get('저장형식'), ('PDF','원본'), '모두'))
                identity = (client.biz, p)
                if identity in seen:
                    raise ValueError('동일 사업자·과세기간이 중복됩니다. 조회항목을 한 행에 통합하세요.')
                seen.add(identity)
                clients.append(client)
            except (ValueError, TypeError) as exc:
                raise ValueError(f'{rownum}행: {exc}') from exc
        if not clients:
            raise ValueError('실행할 거래처가 없습니다.')
        return clients
    finally:
        w.close()


def build_jobs(clients: list[Client]) -> list[Job]:
    jobs = []
    for c in clients:
        p = c.period
        base = {'사업자등록번호':c.biz, '과세연도':str(p.year)}
        for service in c.services:
            if service == 1:
                jobs.append(Job(c,1,'전체보고서',base | {'기수':f'{p.term}기','신고구분':p.filing}))
            elif service == 2:
                for typ in c.invoice_types:
                    for direction in c.directions:
                        jobs.append(Job(c,2,f'{typ}_{direction}',base | {'분류':typ,'구분':direction,
                            '작성일자':'일자별','시작일':p.start,'종료일':p.end}))
            elif service == 3:
                jobs.append(Job(c,3,'매출',base | {'시작분기':f'{min(p.quarters)}분기','종료분기':f'{max(p.quarters)}분기'}))
            elif service == 4:
                # 화면은 조회연도 + 월 선택. 전체를 선택하면 과세기간 외 자료가 섞입니다.
                for month in p.months:
                    jobs.append(Job(c,4,f'{month:02d}월',base | {'조회월':f'{month}월'}))
            elif service == 5:
                for tab in c.special_tabs:
                    for quarter in p.quarters:
                        jobs.append(Job(c,5,f'{tab}_{quarter}분기',base | {'특례조회탭':tab,
                            '적용분기':f'{quarter}분기','품목구분':c.goods,'매입매출구분':c.special_direction}))
            elif service == 6:
                scope = '반기' if p.filing == '확정' and p.scope == '반기' else '분기'
                selection = ('상반기' if p.term==1 else '하반기') if scope=='반기' else f'{p.quarters[0]}분기'
                jobs.append(Job(c,6,'수출실적',base | {'조회단위':scope,'조회기간':selection}))
            elif service == 7:
                if p.filing == '확정' and p.scope == '분기':
                    # 사이트 구매공급일의 확정 구분은 반기이며 4~6/10~12월만 따로 요청할 수 없습니다.
                    raise ValueError(f'{c.row}행: 7번은 분기만 조회하는 확정신고를 지원하지 않습니다. 조회범위=반기로 하거나 7번을 제외하세요.')
                for doc in c.documents:
                    jobs.append(Job(c,7,doc,base | {'서류구분':doc,'기수':f'{p.term}기','신고구분':p.filing}))
    return jobs


def validate_artifact(path: Path) -> dict:
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError('다운로드 파일이 없거나 비어 있습니다.')
    head = path.read_bytes()[:1024].lstrip()
    if head.lower().startswith((b'<!doctype html', b'<html')):
        raise ValueError('자료 대신 HTML/로그인 페이지가 다운로드되었습니다.')
    suffix = path.suffix.lower()
    info = {'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
    if suffix == '.pdf':
        if not head.startswith(b'%PDF-'):
            raise ValueError('PDF 헤더가 없습니다.')
        from pypdf import PdfReader
        reader = PdfReader(str(path), strict=True)
        if reader.is_encrypted or not len(reader.pages):
            raise ValueError('PDF가 암호화되었거나 페이지가 없습니다.')
        info['pages'] = len(reader.pages)
        # 이미지 PDF에는 텍스트가 없을 수 있어 텍스트 일치로 성공을 꾸미지 않습니다.
        info['text_present'] = bool(''.join(p.extract_text() or '' for p in reader.pages).strip())
    elif suffix == '.xlsx':
        wb = load_workbook(path, read_only=True)
        try:
            if not wb.sheetnames:
                raise ValueError('엑셀 시트가 없습니다.')
        finally:
            wb.close()
    elif suffix == '.xls':
        if not (head.startswith(bytes.fromhex('D0CF11E0A1B11AE1')) or b'<table' in head.lower() or b'<xml' in head.lower()):
            raise ValueError('알려진 XLS 형식이 아닙니다.')
        info['structure_checked'] = 'header_only'
    elif suffix not in ('.csv','.txt','.zip'):
        raise ValueError(f'예상하지 못한 확장자: {suffix}')
    elif suffix == '.zip':
        import zipfile
        with zipfile.ZipFile(path) as z:
            if not z.namelist() or z.testzip() is not None:
                raise ValueError('손상/빈 ZIP')
    return info


class Journal:
    """증분 JSONL. 개인정보가 포함되는 로컬 파일이며 외부 전송하지 않습니다."""
    def __init__(self, path: Path):
        self.path = path
        self.latest: dict[str, dict] = {}
        if path.exists():
            lines = path.read_text(encoding='utf-8').splitlines()
            for n, line in enumerate(lines,1):
                try:
                    item = json.loads(line)
                except json.JSONDecodeError:
                    if n == len(lines):  # 중단된 마지막 줄은 별도 보관 후 제거
                        backup = path.with_name(path.name + '.incomplete-' + datetime.now().strftime('%Y%m%d%H%M%S%f'))
                        backup.write_bytes(path.read_bytes())
                        path.write_text('\n'.join(lines[:-1]) + ('\n' if lines[:-1] else ''), encoding='utf-8')
                        continue
                    raise ValueError(f'실행기록 {n}행이 손상되었습니다.')
                self.latest[item['job_key']] = item

    def reusable(self, job: Job) -> bool:
        item = self.latest.get(job.key, {})
        if item.get('status') != 'saved' or not item.get('artifacts'):
            return False
        if not item.get('confirmed'):
            return False
        for a in item['artifacts']:
            try:
                actual = validate_artifact(Path(a['path']))
            except Exception:
                return False
            if actual['sha256'] != a['sha256']:
                return False
        return True

    def record(self, job: Job, status: str, artifacts: list[dict], confirmed: bool, note: str = '') -> None:
        item = {'timestamp':datetime.now().astimezone().isoformat(), 'job_key':job.key,
                'row':job.client.row, 'service':job.service, 'variant':job.variant,
                'status':status,'confirmed':confirmed,'artifacts':artifacts,'note':note}
        self.path.parent.mkdir(parents=True, exist_ok=True)
        # 미완성 줄이 남아 있을 때 다음 JSON을 같은 줄에 붙이지 않습니다.
        if self.path.exists() and self.path.stat().st_size:
            with self.path.open('rb') as f:
                f.seek(-1,2); tail=f.read(1)
            if tail != b'\n':
                with self.path.open('ab') as f:f.write(b'\n')
        with self.path.open('a',encoding='utf-8') as f:
            f.write(json.dumps(item,ensure_ascii=False)+'\n')
        self.latest[job.key] = item
